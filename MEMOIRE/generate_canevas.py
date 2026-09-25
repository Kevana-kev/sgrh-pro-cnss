"""Génère le canevas du mémoire BONGA KASUSA REBECCA."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TITLE = (
    "Conception et implémentation d'un système d'information intelligent "
    "de gestion des ressources humaines intégrant la biométrie (ZK-9500) "
    "et les cartes RFID pour le suivi et l'optimisation du personnel : cas de la CNSS"
)

AUTHOR = "BONGA KASUSA REBECCA"
OUTPUT = r"b:\LARIS\KEVANA\FINAL\bong\MEMOIRE\CANEVAS DU TRAVAIL BONGA KASUSA REBECCA.docx"


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    doc.add_paragraph(text, style=style)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # Page de titre interne
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("CANEVAS DU TRAVAIL")
    run.bold = True
    run.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run(AUTHOR)
    r2.bold = True
    r2.font.size = Pt(14)

    doc.add_paragraph()
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p_title.add_run(TITLE)
    r3.italic = True
    r3.font.size = Pt(11)

    doc.add_page_break()

    # Pages préliminaires
    prelim = [
        "EPIGRAPHE",
        "IN MEMORIAM",
        "DÉDICACE",
        "REMERCIEMENTS",
        "ABRÉVIATIONS ET SIGLES",
        "LISTE DES FIGURES",
        "LISTE DES GRAPHIQUES",
        "LISTE DES TABLEAUX",
        "RÉSUMÉ (FRANÇAIS)",
        "ABSTRACT (RÉSUMÉ EN ANGLAIS)",
        "RÉSUMÉ (LINGALA)",
    ]
    add_heading(doc, "PAGES PRÉLIMINAIRES", level=1)
    for item in prelim:
        doc.add_paragraph(item, style="List Number")

    doc.add_page_break()

    # Introduction
    add_heading(doc, "INTRODUCTION GÉNÉRALE", level=1)
    intro_sections = [
        ("1. Contexte et justification", [
            "Présentation de la CNSS en RDC (siège central)",
            "Importance de la gestion des ressources humaines dans les institutions publiques",
            "Constats sur le terrain (stage) : registres papier, Excel, ancien logiciel, pointage manuel",
            "Problèmes identifiés : fraudes, lenteurs administratives, absence de centralisation, faible traçabilité",
            "Nécessité d'un SIGRH moderne intégrant biométrie et RFID",
        ]),
        ("2. Problématique et hypothèses", [
            "2.1. Problématique",
            "2.2. Questions de recherche (principale et spécifiques)",
            "2.3. Hypothèses (principale et spécifiques)",
        ]),
        ("3. Objectifs général et spécifiques", [
            "3.1. Objectif général",
            "3.2. Objectifs spécifiques (centralisation, automatisation, biométrie/RFID, KPI, espace agent)",
        ]),
        ("4. Choix et intérêt du sujet", [
            "4.1. Choix du sujet",
            "4.2. Intérêt du sujet (académique, institutionnel, pratique)",
        ]),
        ("5. Revue de la littérature (aperçu)", []),
        ("6. Méthodologie de recherche", [
            "Approche méthodologique (analyse, conception, développement, validation)",
            "Techniques : entretiens, observation, analyse documentaire, prototypage",
            "Outils : UML (cas d'utilisation, classes, séquences, déploiement), Laravel, base de données",
        ]),
        ("7. Délimitation du sujet", [
            "Périmètre thématique : GRH interne CNSS (hors gestion des assurés/cotisations)",
            "Périmètre technique : prototype hébergeable, biométrie ZK-9500 + RFID",
            "Périmètre géographique : siège central CNSS (RDC)",
            "Périmètre temporel : période du mémoire",
        ]),
        ("8. Subdivision du travail", []),
    ]
    for title, subs in intro_sections:
        add_heading(doc, title, level=2)
        for s in subs:
            add_bullet(doc, s)

    doc.add_page_break()

    # Chapitre I
    add_heading(doc, "CHAPITRE I : CADRE THÉORIQUE ET ÉTAT DE L'ART", level=1)
    ch1 = [
        ("I.1. Introduction du chapitre", []),
        ("I.2. Généralités sur la gestion des ressources humaines (GRH)", [
            "Définition et enjeux de la GRH",
            "Fonctions RH : recrutement, administration, paie, formation, évaluation",
            "GRH dans les institutions publiques africaines",
        ]),
        ("I.3. Systèmes d'information des ressources humaines (SIGRH / SIRH)", [
            "Définition et composantes d'un SIGRH",
            "Avantages d'un SIGRH intégré",
            "Architecture fonctionnelle d'un SIGRH",
        ]),
        ("I.4. Dimension « intelligente » du système", [
            "Tableaux de bord et indicateurs clés (KPI RH)",
            "Alertes automatiques (absences, retards, fin de contrat, budgets)",
            "Aide à la décision et optimisation du personnel",
            "Différence entre automatisation et intelligence décisionnelle",
        ]),
        ("I.5. Technologies d'identification et de contrôle d'accès", [
            "Biométrie : principes, types (empreinte digitale), avantages et limites",
            "Lecteur ZK-Teco 9500 : caractéristiques et usages",
            "Cartes RFID : fonctionnement, attribution, sécurité",
            "Combinaison empreinte + RFID dans un contexte RH",
            "Aspects éthiques et protection des données biométriques (prototype)",
        ]),
        ("I.6. Revue des solutions existantes", [
            "Solutions internationales (SAP SuccessFactors, Oracle HCM, Workday)",
            "Solutions open source et PME (Odoo RH, Dolibarr, etc.)",
            "Solutions et pratiques en Afrique / RDC",
            "Comparatif fonctionnel et technologique",
        ]),
        ("I.7. Limites des systèmes traditionnels de gestion RH", [
            "Gestion manuelle et outils bureautiques (Excel, registres)",
            "Absence d'intégration et de traçabilité",
            "Faible fiabilité du pointage",
        ]),
        ("I.8. Modélisation UML et architectures logicielles", [
            "Intérêt de l'UML dans la conception de SI",
            "Diagrammes retenus (cas d'utilisation, classes, séquences, activités, déploiement)",
            "Architecture n-tiers et séparation backend / frontend",
        ]),
        ("I.9. Conclusion du chapitre", []),
    ]
    for title, subs in ch1:
        add_heading(doc, title, level=2)
        for s in subs:
            add_bullet(doc, s)

    doc.add_page_break()

    # Chapitre II
    add_heading(doc, "CHAPITRE II : ANALYSE DU SYSTÈME ACTUEL ET SPÉCIFICATION DES BESOINS", level=1)
    ch2 = [
        ("II.1. Introduction du chapitre", []),
        ("II.2. Présentation institutionnelle de la CNSS (RDC)", [
            "Historique et missions de la CNSS",
            "Organisation du siège central",
            "Organigramme et services RH",
            "Effectif et typologie du personnel",
        ]),
        ("II.3. Analyse du système actuel de gestion RH", [
            "Processus actuels de gestion du personnel",
            "Outils utilisés (papier, Excel, ancien logiciel)",
            "Gestion actuelle des présences et pointage manuel",
            "Gestion des congés, transmission administrative et contrats",
            "Problèmes et dysfonctionnements identifiés",
            "Analyse des fraudes et irrégularités de pointage",
            "Impacts sur la productivité et la prise de décision",
        ]),
        ("II.4. Analyse des besoins des acteurs", [
            "Besoins des administrateurs / RH",
            "Besoins des managers et superviseurs",
            "Besoins des agents (espace personnel)",
            "Besoins en identification fiable (biométrie + RFID)",
            "Besoins en reporting et tableaux de bord",
        ]),
        ("II.5. Scénarios d'utilisation", [
            "Enrôlement biométrique et attribution carte RFID",
            "Pointage entrée/sortie par empreinte digitale",
            "Demande et validation de congé (Agent → Manager → RH)",
            "Consolidation administrative et export vers la finance",
            "Consultation KPI par le responsable RH",
        ]),
        ("II.6. Spécification fonctionnelle — les 8 modules", [
            "Module 1 : Gestion des utilisateurs et accès (RBAC, audit)",
            "Module 2 : Gestion des employés (dossier RH)",
            "Module 3 : Intégration biométrique (ZK-9500 + RFID)",
            "Module 4 : Gestion des présences et absences",
            "Module 5 : Gestion des congés et permissions",
            "Module 6 : Suivi administratif des éléments de rémunération",
            "Module 7 : Évaluation des performances",
            "Module 8 : Tableaux de bord et rapports (KPI)",
        ]),
        ("II.7. Spécification technique et contraintes", [
            "Architecture backend / frontend séparée",
            "Base de données relationnelle (hébergement cloud/local)",
            "API REST sécurisée (Sanctum, bcrypt)",
            "Bridge biométrique C# (ZK-Teco SDK, port 5002)",
            "Interface web responsive pour RH et agents",
            "Exigences de sécurité, performance et disponibilité",
        ]),
        ("II.8. Cahier des charges fonctionnel et technique", []),
        ("II.9. Cadrage et planification du projet (PERT, Gantt, budget)", []),
        ("II.10. Conclusion du chapitre", []),
    ]
    for title, subs in ch2:
        add_heading(doc, title, level=2)
        for s in subs:
            add_bullet(doc, s)

    doc.add_page_break()

    # Chapitre III
    add_heading(doc, "CHAPITRE III : CONCEPTION ET ARCHITECTURE DU SYSTÈME", level=1)
    ch3 = [
        ("III.1. Introduction du chapitre", []),
        ("III.2. Architecture globale du système", [
            "Vue d'ensemble (client web, API backend, base de données, bridge biométrique, périphériques)",
            "Schéma d'architecture logique et physique",
            "Flux de données entre composants",
        ]),
        ("III.3. Modélisation UML", [
            "Diagramme de cas d'utilisation (Admin RH, Manager, Agent, Système biométrique)",
            "Diagramme de classes (Employé, User, Attendance, RemunerationElement, Leave, Contract, etc.)",
            "Diagrammes de séquence (pointage biométrique, validation congé, consolidation administrative)",
            "Diagramme d'activités (workflow congé, processus de pointage)",
            "Diagramme de déploiement (PC responsable RH, serveur, lecteur ZK-9500, cartes RFID)",
        ]),
        ("III.4. Conception de la base de données", [
            "Modèle entité-relation (MER)",
            "Description des tables principales",
            "Relations et contraintes d'intégrité",
            "Gestion des données biométriques (templates) et RFID",
        ]),
        ("III.5. Conception du backend (API REST)", [
            "Structure en couches (routes, services, modèles, middlewares)",
            "Endpoints par module (auth, employés, éléments de rémunération, présences, congés, biométrie, rapports)",
            "Sécurité : Sanctum, RBAC, journal d'audit",
            "Choix technologiques (Laravel, Eloquent, validation)",
        ]),
        ("III.6. Conception du frontend (interface web)", [
            "Structure des écrans (dashboard RH personnalisable, terminal de pointage, espace agent)",
            "Tableau de bord KPI et graphiques",
            "Module biométrie (enrôlement, statut bridge, gestion RFID)",
            "Ergonomie et charte visuelle",
        ]),
        ("III.7. Conception du sous-système biométrique", [
            "Architecture du Fingerprint Bridge (C# headless)",
            "Intégration SDK ZKFinger (ZK-9500)",
            "Protocole bridge ↔ backend (/scan, /match, port 5002)",
            "Gestion des cartes RFID (attribution, activation, désactivation)",
            "Processus d'enrôlement et pointage entrée/sortie",
            "Installateur Windows one-touch et sécurisation (clé API, DPAPI)",
        ]),
        ("III.8. Conception des modules métier RH", [
            "M1 Accès et M2 Employés (RBAC, dossier RH)",
            "M3 Biométrie et M4 Présences (pointage, retards)",
            "M5 Congés (workflow de validation, soldes)",
            "M6 Éléments de rémunération (états administratifs, export finance)",
            "M7 Évaluation et M8 Rapports / KPI",
        ]),
        ("III.9. Conception de la dimension intelligente", [
            "Indicateurs KPI (effectifs, absentéisme, ponctualité, évaluations)",
            "Moteur d'alertes (retards, absences, congés en attente)",
            "Tableaux de bord par profil (RH, Manager, Agent) et graphiques personnalisables",
        ]),
        ("III.10. Stratégie d'hébergement et déploiement", [
            "Hébergement mutualisé PHP/MySQL et bridge local Windows",
            "Environnements (développement SQLite, production MySQL)",
            "Prototypage sur PC du responsable RH (démonstration)",
        ]),
        ("III.11. Conclusion du chapitre", []),
    ]
    for title, subs in ch3:
        add_heading(doc, title, level=2)
        for s in subs:
            add_bullet(doc, s)

    doc.add_page_break()

    # Chapitre IV
    add_heading(doc, "CHAPITRE IV : IMPLÉMENTATION, TESTS ET VALIDATION", level=1)
    ch4 = [
        ("IV.1. Introduction du chapitre", []),
        ("IV.2. Environnement de développement et outils", [
            "Matériel et logiciels utilisés",
            "Stack technique retenue (PHP/Laravel, HTML/CSS/JS, SQLite/MySQL, C# bridge)",
            "Périphériques : ZK-Teco 9500, cartes RFID, lecteur RFID",
        ]),
        ("IV.3. Implémentation du backend", [
            "Mise en place de l'API REST",
            "Modèles de données et migrations",
            "Authentification Sanctum et gestion RBAC",
            "Implémentation des 8 modules métier",
        ]),
        ("IV.4. Implémentation du frontend", [
            "Interface de connexion et gestion des sessions",
            "Dashboard RH avec KPI animés et graphiques",
            "Espace agent (présences, heures, retards, évaluations, congés)",
            "Formulaires CRUD (employés, départements, contrats, etc.)",
        ]),
        ("IV.5. Implémentation du sous-système biométrique", [
            "Développement du Fingerprint Bridge et installateur Windows one-touch",
            "Intégration SDK ZK-Teco 9500",
            "Enrôlement des empreintes digitales",
            "Attribution et gestion des cartes RFID",
            "Pointage réel par empreinte (entrée/sortie)",
            "Synchronisation avec le backend RH",
        ]),
        ("IV.6. Implémentation des fonctionnalités intelligentes", [
            "Calcul automatique des KPI RH",
            "Génération des rapports (PDF, Excel, CSV)",
            "Système d'alertes et notifications",
            "Statistiques par département et évolution temporelle",
        ]),
        ("IV.7. Intégration et déploiement du prototype", [
            "Assemblage backend + frontend + bridge biométrique",
            "Configuration sur PC du responsable RH",
            "Jeux de données de test / données réelles anonymisées",
        ]),
        ("IV.8. Tests et validation", [
            "Tests unitaires et tests d'intégration",
            "Tests fonctionnels par module (8 modules)",
            "Tests du pointage biométrique et RFID",
            "Tests de sécurité (authentification, permissions)",
            "Tests de performance et ergonomie",
            "Scénarios de validation (cas d'utilisation du chapitre II)",
        ]),
        ("IV.9. Résultats et analyse", [
            "Présentation des interfaces réalisées (captures d'écran)",
            "Résultats des tests de pointage biométrique",
            "Comparaison avant/après (système manuel vs prototype)",
            "Analyse des KPI produits",
            "Retours utilisateurs (RH, agents) — stage CNSS",
        ]),
        ("IV.10. Limites du prototype et perspectives", [
            "Limites techniques (un seul poste de pointage, SQLite local vs MySQL)",
            "Limites fonctionnelles et juridiques (données biométriques)",
            "Perspectives : déploiement multi-sites, mobile, IA avancée",
        ]),
        ("IV.11. Conclusion du chapitre", []),
    ]
    for title, subs in ch4:
        add_heading(doc, title, level=2)
        for s in subs:
            add_bullet(doc, s)

    doc.add_page_break()

    # Conclusion générale
    add_heading(doc, "CONCLUSION GÉNÉRALE", level=1)
    conclusion = [
        "Synthèse des travaux réalisés",
        "Rappel de la problématique et des objectifs",
        "Validation des hypothèses",
        "Apports du système pour la CNSS (siège central)",
        "Apports académiques et personnels",
        "Limites de l'étude",
        "Perspectives d'évolution (hébergement, déploiement réel, modules avancés)",
    ]
    for c in conclusion:
        add_bullet(doc, c)

    doc.add_page_break()

    # Bibliographie et annexes
    add_heading(doc, "BIBLIOGRAPHIE", level=1)
    doc.add_paragraph("(À compléter : ouvrages GRH, articles SIGRH, documentation ZK-Teco, normes UML, sites CNSS RDC)")

    add_heading(doc, "ANNEXES", level=1)
    annexes = [
        "Annexe A : Cahier des charges complet",
        "Annexe B : Diagrammes UML complets",
        "Annexe C : Maquettes des interfaces",
        "Annexe D : Extraits de code significatifs",
        "Annexe E : Guides d'utilisation (enrôlement biométrique, pointage RFID)",
        "Annexe F : Questionnaire / grille d'entretien (RH, agents)",
        "Annexe G : Captures d'écran du système implémenté",
        "Annexe H : Tableau comparatif des solutions SIGRH",
    ]
    for a in annexes:
        add_bullet(doc, a)

    doc.save(OUTPUT)
    print(f"Canevas créé : {OUTPUT}")


if __name__ == "__main__":
    build()
